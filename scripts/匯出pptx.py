# -*- coding: utf-8 -*-
"""把成語圖卡版面產成 PPTX（16:9），可直接投影，也可匯入 Canva 再編輯。

用法: python3 scripts/匯出pptx.py
輸出: output/成語卡.pptx（每個成語一頁，文字框都是可編輯的真文字框）
"""
import json, os
import 標籤 as 標籤模組
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data", "成語資料.json")
OUT  = os.path.join(ROOT, "output", "成語卡.pptx")

FONT   = "Noto Sans TC"          # Canva 內建，繁體字形正確
PX     = 6350                    # 1 px → EMU（1920px = 13.333in）
BG     = "FAF6EE"; INK = "2B2B2B"; RED = "8C2F26"
GOLD   = "C9A227"; GRAY = "6E6459"
WARNBG = "FBEFE4"; WARNFG = "B4501F"; LINE = "DCD2C0"; WARNLINE = "E8C9AE"

px = lambda v: Emu(int(v * PX))
sz = lambda v: Pt(v * 0.5)       # 1920px 寬 = 960pt，故 1px = 0.5pt

def set_cjk(run, name):
    """PPTX 需另外指定東亞字型，中文才會套用"""
    rPr = run.font._element
    for tag in ("a:ea", "a:cs"):
        e = rPr.makeelement(qn(tag), {"typeface": name})
        rPr.append(e)

def text(slide, x, y, w, h, s, size, color, bold=False,
         align=PP_ALIGN.LEFT, spacing=None, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.2
    r = p.add_run(); r.text = s
    r.font.size = sz(size); r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = FONT
    set_cjk(r, FONT)
    if spacing:                                  # 字距（1/100 pt）
        r.font._element.set("spc", str(int(spacing * 0.5 * 100)))
    return tb

def shape(slide, kind, x, y, w, h, fill, line=None, line_w=2):
    sp = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if fill:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = px(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    sp.text_frame.text = ""
    return sp

def build(prs, item):
    s = prs.slides.add_slide(prs.slide_layouts[6])      # 空白版面
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)

    shape(s, MSO_SHAPE.RECTANGLE, 26, 26, 1868, 1028, None, LINE, 3)

    # 課次標籤
    tag = item["標籤"]
    shape(s, MSO_SHAPE.RECTANGLE, 96, 74, 358, 54, RED)
    text(s, 118, 78, 320, 46, tag, 32, "FFF8F0", True)

    # 成語區
    zy = item["注音"].split()                      # 每組注音對準下方的字
    for i, z in enumerate(zy):
        text(s, 540 + i * 210, 186, 210, 46, z, 34, GRAY, align=PP_ALIGN.CENTER)
    text(s, 540, 236, 840, 200, item["成語"], 156, INK, True,
         align=PP_ALIGN.CENTER, spacing=54)
    text(s, 540, 414, 840, 44, item.get("拼音", ""), 30, GOLD, align=PP_ALIGN.CENTER)
    shape(s, MSO_SHAPE.RECTANGLE, 96, 468, 1728, 2, LINE)

    # 解釋
    shape(s, MSO_SHAPE.RECTANGLE, 96, 514, 8, 34, GOLD)
    text(s, 120, 508, 200, 50, "解釋", 38, RED, True)
    text(s, 120, 568, 1704, 60, item["解釋"], 34, INK)
    text(s, 120, 624, 1704, 50, "（白話）" + item["白話解釋"], 30, GRAY)

    # 例句
    shape(s, MSO_SHAPE.RECTANGLE, 96, 700, 8, 34, GOLD)
    text(s, 120, 694, 200, 50, "例句", 38, RED, True)
    for i, ex in enumerate(item["例句"]):
        y = 754 + i * 60
        shape(s, MSO_SHAPE.OVAL, 120, y + 4, 34, 34, RED)
        text(s, 120, y + 7, 34, 34, str(i + 1), 30, "FFF8F0", True,
             align=PP_ALIGN.CENTER)
        text(s, 174, y, 1650, 56, ex, 34, INK)

    # 易錯提醒
    if item.get("易錯提醒"):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 96, 920, 1728, 86, WARNBG, WARNLINE, 2)
        text(s, 122, 943, 40, 44, "※", 30, WARNFG)
        text(s, 166, 943, 1598, 60, item["易錯提醒"], 28, WARNFG)

def main():
    db = 標籤模組.標註(json.load(open(DATA, encoding="utf-8")))
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    for item in db["成語"]:
        build(prs, item)
    prs.save(OUT)
    print(f"已產出 {len(db['成語'])} 頁 → {OUT}")

if __name__ == "__main__":
    main()
